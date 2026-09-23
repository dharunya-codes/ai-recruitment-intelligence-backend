from __future__ import annotations

import re
from pathlib import Path

import fitz
from docx import Document


def extract_text_from_pdf(path: str | Path) -> str:
    with fitz.open(path) as document:
        pages = [page.get_text("text") for page in document]
    return "\n".join(pages).strip()


def extract_text_from_docx(path: str | Path) -> str:
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_candidate_info(text: str) -> dict[str, str | None]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    email_match = re.search(r"\b[\w.+-]+@[\w-]+(?:\.[\w-]+)+\b", text)
    phone_match = re.search(
        r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)|"
        r"(?<!\d)\+?[1-9]\d{1,3}[\s.-]?(?:\(\d{2,4}\)[\s.-]?)?\d[\d\s.-]{6,}\d(?!\d)",
        text,
    )

    name = None
    for line in lines[:8]:
        if "@" in line or re.search(r"\d", line):
            continue
        words = line.split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z][A-Za-z.'-]*", word) for word in words):
            name = line
            break

    phone = phone_match.group(0).strip() if phone_match else None
    return {
        "name": name or "Unknown Candidate",
        "email": email_match.group(0) if email_match else None,
        "phone": phone,
    }